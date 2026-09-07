"""Un extractor de texto por tipo de archivo. Cada funcion toma un Path y
devuelve el texto plano del documento, con marcadores simples de posicion
(pagina/hoja/diapositiva) para que quien use este texto mas adelante (el
RAG de la Fase 6) pueda seguir sabiendo de que parte del documento salio
cada fragmento, sin tener que volver a abrir el archivo original.

python-docx/openpyxl/python-pptx ya eran dependencias de Atlas (para
*escribir* documentos, ver skills/documents.py) - la misma libreria sirve
para leerlos, no hizo falta sumar nada nuevo para esos tres formatos."""

import csv
from pathlib import Path

import fitz  # PyMuPDF - solo se usa para renderizar paginas sin texto a imagen
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from core.documents.ocr import ocr_image


def _ocr_pdf_page(fitz_doc: "fitz.Document", page_index: int, zoom: float = 2.0) -> str:
    # zoom 2.0 ~= 144 DPI (el render por defecto es ~72 DPI) - Tesseract
    # rinde notablemente mejor con mas resolucion que con el tamaño nativo.
    page = fitz_doc[page_index]
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    return ocr_image(image)


def extract_pdf(path: Path) -> str:
    """Extrae el texto de cada pagina; si una pagina no tiene capa de texto
    (tipico de un escaneo), esa pagina puntual se renderiza a imagen y se
    le hace OCR - asi un PDF mixto (algunas paginas con texto real, otras
    escaneadas) no pierde las paginas que si tenian texto nativo, y el
    costo de OCR (mucho mas lento que leer texto embebido) solo se paga en
    las paginas que de verdad lo necesitan."""
    reader = PdfReader(str(path))
    fitz_doc = None
    ocr_unavailable_reason = None
    pages = []
    try:
        for i, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                if fitz_doc is None:
                    fitz_doc = fitz.open(str(path))
                try:
                    text = _ocr_pdf_page(fitz_doc, i - 1)
                except RuntimeError as exc:
                    ocr_unavailable_reason = str(exc)
                    text = ""
            if text:
                pages.append(f"--- Página {i} ---\n{text}")
    finally:
        if fitz_doc is not None:
            fitz_doc.close()

    if not pages:
        if ocr_unavailable_reason:
            raise ValueError(
                f"el PDF parece un escaneo (sin texto embebido) y {ocr_unavailable_reason}"
            )
        raise ValueError(
            "no se pudo extraer texto de este PDF, ni siquiera con OCR "
            "(¿está vacío o la imagen es ilegible?)."
        )
    return "\n\n".join(pages)


def extract_image(path: Path) -> str:
    """OCR directo sobre una imagen suelta (foto de un documento, un
    recibo escaneado, etc.) - no interpreta la imagen visualmente (para
    eso está la skill see_screen con Gemini Vision), solo extrae el texto
    legible que contenga."""
    image = Image.open(path)
    text = ocr_image(image)
    if not text:
        raise ValueError("no se encontró texto legible en la imagen.")
    return text


def extract_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("el documento Word no tiene texto (¿está vacío?).")
    return text


def extract_xlsx(path: Path) -> str:
    wb = load_workbook(str(path), data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                rows.append(" | ".join(values))
        if rows:
            sheets.append(f"--- Hoja: {ws.title} ---\n" + "\n".join(rows))

    if not sheets:
        raise ValueError("el Excel no tiene datos (¿está vacío?).")
    return "\n\n".join(sheets)


def extract_csv(path: Path) -> str:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        rows = [" | ".join(row) for row in csv.reader(f) if any(cell.strip() for cell in row)]
    if not rows:
        raise ValueError("el CSV no tiene datos (¿está vacío?).")
    return "\n".join(rows)


def extract_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
        if lines:
            slides.append(f"--- Diapositiva {i} ---\n" + "\n".join(lines))

    if not slides:
        raise ValueError("la presentación no tiene texto (¿está vacía?).")
    return "\n\n".join(slides)


def extract_html(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        raise ValueError("la página HTML no tiene texto visible.")
    return "\n".join(lines)
